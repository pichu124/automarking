"""Extract plain text from uploaded documents.

Supports .docx, .pdf, .txt/.md, and .zip archives (which are expanded into one
entry per contained document). Everything returned is plain text keyed by a
display name, which downstream code feeds to the LLM.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from typing import List

import docx  # python-docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree  # installed with python-docx
from pypdf import PdfReader


_TEXT_EXTS = (".txt", ".md", ".csv")
# Archiver junk to skip inside zips. Hidden dotfiles (.DS_Store) are matched on
# the BASENAME where members are read — a leading "./" or "../" on an otherwise
# normal path is not junk (tar-to-zip converters emit those) and must not cause
# the member to be silently dropped.
_SKIP_PREFIXES = ("__MACOSX/",)
# Zip-bomb guards: uploads are untrusted and the server is shared — a 284 KB
# zip can hold a 300 MB file, and a 42.zip-style archive nests without end.
_MAX_MEMBER_BYTES = 50 * 1024 * 1024       # one decompressed file in a zip
_MAX_ZIP_TOTAL_BYTES = 300 * 1024 * 1024   # all of one upload, nested included
_MAX_ZIP_DEPTH = 2                         # LMS export zip -> a student's zip


@dataclass
class ParsedDocument:
    name: str
    text: str


def _read_docx(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        # python-docx refuses slightly-broken packages outright — e.g. a
        # dangling relationship to a deleted image ("There is no item named
        # 'word/media/imageN.png' in the archive") after a bad edit/repair.
        # The same goes for XML with raw control characters. The text itself
        # is intact in word/document.xml, so read the XML parts straight from
        # the zip instead of failing the submission.
        return _docx_xml_text(data)
    parts = [
        _parse_docx_xml(rel.target_part.blob)
        for rel in document.part.rels.values()
        if rel.reltype in (RT.FOOTNOTES, RT.ENDNOTES, RT.HEADER, RT.FOOTER)
    ]
    return _docx_text(document.element.body, parts)


def _docx_xml_text(data: bytes) -> str:
    """Same extraction as _read_docx, locating the parts by their usual names."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        body = _parse_docx_xml(zf.read("word/document.xml"))
        parts = [
            _parse_docx_xml(zf.read(n)) for n in zf.namelist()
            if re.fullmatch(r"word/(footnotes|endnotes|header\d*|footer\d*)\.xml", n)
        ]
    return _docx_text(body, parts)


# recover=True drops characters XML forbids instead of rejecting the whole
# part; entities stay unresolved because the file is untrusted upload data.
_DOCX_XML_PARSER = etree.XMLParser(recover=True, resolve_entities=False)


def _parse_docx_xml(blob: bytes):
    return etree.fromstring(blob, _DOCX_XML_PARSER)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_M_T = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"
# Never text: paragraph properties (their w:tab elements are tab stops), the
# VML copy Word writes of every text box (mc:Fallback), tracked deletions and
# moved-away text. Field codes (w:instrText) and deleted runs (w:delText) are
# skipped by simply not being w:t.
_DOCX_SKIP = frozenset({
    "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback",
    _W + "pPr", _W + "del", _W + "moveFrom",
})
_DOCX_CHARS = {
    _W + "tab": "\t", _W + "ptab": "\t", _W + "br": "\n", _W + "cr": "\n",
    _W + "noBreakHyphen": "-",
}
_NOTE_SEPARATORS = frozenset({"separator", "continuationSeparator", "continuationNotice"})
# Parts appended after the body, in this order, each under its heading.
_DOCX_EXTRA_PARTS = (
    ("[Footnotes]", (_W + "footnotes",)),
    ("[Endnotes]", (_W + "endnotes",)),
    ("[Headers and footers]", (_W + "hdr", _W + "ftr")),
)


def _docx_text(body, parts) -> str:
    """Body text in document order, then notes and headers/footers."""
    lines = _docx_blocks(body)
    for heading, tags in _DOCX_EXTRA_PARTS:
        extra = [line for root in parts if root.tag in tags for line in _docx_blocks(root)]
        if extra:
            lines += ["", heading, *extra]
    return "\n".join(lines).strip()


def _docx_blocks(element) -> List[str]:
    """Text lines under `element` in document order: one per non-empty
    paragraph, one per table row (cells joined by " | "). Descends through
    content controls, text boxes and any other wrapper that holds paragraphs."""
    blocks: List[str] = []
    for child in element:
        if child.tag in _DOCX_SKIP or child.get(_W + "type") in _NOTE_SEPARATORS:
            continue
        if child.tag == _W + "p":
            blocks.extend(_docx_paragraph(child))
        elif child.tag == _W + "tr":
            # Cells may sit inside a content control; nested tables' cells don't count.
            cells = [
                "\n".join(_docx_blocks(tc)).strip() for tc in child.iter(_W + "tc")
                if next(tc.iterancestors(_W + "tr")) is child
            ]
            if any(cells):
                blocks.append(" | ".join(cells))
        else:
            blocks.extend(_docx_blocks(child))
    return blocks


def _docx_paragraph(paragraph) -> List[str]:
    """The paragraph's own text (runs, content controls, equations), then the
    lines of any text box anchored in it."""
    chars: List[str] = []
    boxes: List[str] = []

    def visit(element) -> None:
        for child in element:
            if child.tag in _DOCX_SKIP:
                continue
            if child.tag in (_W + "t", _M_T):
                chars.append(child.text or "")
            elif child.tag in _DOCX_CHARS:
                chars.append(_DOCX_CHARS[child.tag])
            elif child.tag == _W + "txbxContent":
                boxes.extend(_docx_blocks(child))
            else:
                visit(child)

    visit(paragraph)
    text = "".join(chars)
    return ([text] if text.strip() else []) + boxes


def _read_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts: List[str] = []
    for page in reader.pages:
        # "plain" mode avoids pypdf's layout-mode width maths, which can crash
        # on PDFs whose font-width entries are unresolved indirect objects
        # (TypeError: unsupported operand type(s) for <: 'float' and ...).
        # A per-page guard means one bad page never sinks the whole document.
        try:
            parts.append(page.extract_text(extraction_mode="plain") or "")
        except TypeError:  # older pypdf without the kwarg (already plain)
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                parts.append("")
        except Exception:
            parts.append("")
    return "\n".join(parts)


def _read_text(data: bytes) -> str:
    """Decode text robustly: Excel-exported CSVs are often UTF-8-with-BOM,
    UTF-16 (with or without a BOM) or cp1252 rather than plain UTF-8."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return data.decode("utf-16")
        except UnicodeError:
            pass
    # UTF-16 WITHOUT a BOM decodes as "valid" UTF-8/cp1252 (NUL bytes are legal
    # there), yielding N-U-L-interleaved mojibake. NULs never appear in real
    # text, so their presence means: try UTF-16 first.
    if b"\x00" in data:
        for encoding in ("utf-16-le", "utf-16-be"):
            try:
                candidate = data.decode(encoding)
            except UnicodeError:
                continue
            if "\x00" not in candidate and not _looks_binary(candidate):
                return candidate
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _looks_binary(text: str) -> bool:
    """True when decoded 'text' is really binary garbage (mojibake guard)."""
    sample = text[:4000]
    if not sample:
        return False
    bad = sum(
        1 for ch in sample
        if ch == "�" or (ord(ch) < 32 and ch not in "\t\n\r")
    )
    return bad > len(sample) * 0.10


def _read_xlsx(data: bytes) -> str:
    """Extract text from an Excel workbook, one row per line, sheets labelled.

    Rubrics and grade sheets often live in Excel. ``data_only=True`` returns
    computed cell values instead of formula strings.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - import guard
        raise RuntimeError(
            "The 'openpyxl' package is required to read Excel files. "
            "Install it with: pip install openpyxl"
        ) from exc

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts: List[str] = []
    try:
        for sheet in workbook.worksheets:
            if len(workbook.worksheets) > 1:
                parts.append(f"=== Sheet: {sheet.title} ===")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    parts.append(" | ".join(cells).rstrip(" |"))
    finally:
        workbook.close()
    return "\n".join(parts)


def _read_xlsb(data: bytes) -> str:
    """Extract text from an Excel Binary Workbook (.xlsb) via pyxlsb."""
    try:
        from pyxlsb import open_workbook
    except ImportError as exc:  # pragma: no cover - import guard
        raise RuntimeError(
            "The 'pyxlsb' package is required to read .xlsb files. "
            "Install it with: pip install pyxlsb"
        ) from exc
    parts: List[str] = []
    with open_workbook(io.BytesIO(data)) as wb:
        for sheet_name in wb.sheets:
            with wb.get_sheet(sheet_name) as sheet:
                rows_out: List[str] = []
                for row in sheet.rows():
                    cells = ["" if c.v is None else str(c.v).strip() for c in row]
                    if any(cells):
                        rows_out.append(" | ".join(cells).rstrip(" |"))
                if rows_out:
                    if len(wb.sheets) > 1:
                        parts.append(f"=== Sheet: {sheet_name} ===")
                    parts.extend(rows_out)
    return "\n".join(parts)


def _read_xls(data: bytes) -> str:
    """Extract text from a legacy Excel 97-2003 workbook (.xls) via xlrd."""
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - import guard
        raise RuntimeError(
            "The 'xlrd' package is required to read legacy .xls files. "
            "Install it with: pip install xlrd"
        ) from exc
    book = xlrd.open_workbook(file_contents=data)
    parts: List[str] = []
    for sheet in book.sheets():
        rows_out: List[str] = []
        for ri in range(sheet.nrows):
            cells = [
                "" if v in ("", None) else str(v).strip()
                for v in sheet.row_values(ri)
            ]
            if any(cells):
                rows_out.append(" | ".join(cells).rstrip(" |"))
        if rows_out:
            if book.nsheets > 1:
                parts.append(f"=== Sheet: {sheet.name} ===")
            parts.extend(rows_out)
    return "\n".join(parts)


# Embedded images (screenshots, charts) cannot be text-extracted. Count them
# and prepend a note so the grader never mistakes invisible-for-absent — a
# student's appendix full of rate screenshots must not be called "empty".
_IMAGE_EXTS = (
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif",
    ".emf", ".wmf", ".webp",
)


def _count_opc_images(data: bytes, prefix: str) -> int:
    """Count image files inside an OPC (zip) container under `prefix`.

    Works for Word (word/media/) and every Excel flavour that is a zip —
    .xlsx, .xlsm and .xlsb all store pasted images (e.g. Solver-dialog
    screenshots) as ordinary files under xl/media/.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return sum(
                1 for n in zf.namelist()
                if n.lower().startswith(prefix) and n.lower().endswith(_IMAGE_EXTS)
            )
    except Exception:
        return 0


def _count_docx_images(data: bytes) -> int:
    return _count_opc_images(data, "word/media/")


def _count_xl_images(data: bytes) -> int:
    return _count_opc_images(data, "xl/media/")


def _count_pdf_images(data: bytes) -> int:
    try:
        reader = PdfReader(io.BytesIO(data))
        total = 0
        for page in reader.pages:
            try:
                total += sum(1 for _ in page.images)
            except Exception:
                continue
        return total
    except Exception:
        return 0


def _image_note(count: int) -> str:
    if count <= 0:
        return ""
    return (
        f"[NOTE: this document contains {count} embedded image(s) - e.g. "
        "screenshots, charts, figures - whose content cannot be extracted as "
        "text here. Do NOT state that screenshots, figures or appendix "
        "evidence are missing or absent; they may be in these images. If "
        "required evidence is not visible in the text, say it could not be "
        "verified from the text and grade what is verifiable.]\n\n"
    )


def _embedded_image_note(lowered: str, data: bytes) -> str:
    """The image-count note for any container format that can embed images."""
    if lowered.endswith(".docx"):
        return _image_note(_count_docx_images(data))
    if lowered.endswith((".xlsx", ".xlsm", ".xlsb")):
        return _image_note(_count_xl_images(data))
    if lowered.endswith(".pdf"):
        return _image_note(_count_pdf_images(data))
    return ""


def parse_bytes(name: str, data: bytes) -> List[ParsedDocument]:
    """Parse one uploaded file. A .zip expands to multiple ParsedDocuments."""
    lowered = name.lower()
    if lowered.endswith(".zip"):
        try:
            return _parse_zip(data)
        except Exception as exc:
            # A truncated download or a non-zip renamed .zip raises BadZipFile
            # — surface a loud placeholder instead of crashing the app.
            return [ParsedDocument(name, f"[Could not read {name}: {exc}]")]

    # A standalone image (e.g. a screenshot submitted as its own file) has no
    # extractable text — surface an explicit note instead of binary garbage, so
    # after per-student grouping the grader knows the image evidence EXISTS.
    if lowered.endswith(_IMAGE_EXTS):
        return [ParsedDocument(name, (
            "[Image file - visual content (e.g. a screenshot) that cannot be "
            "extracted as text. Do NOT treat the evidence it may contain as "
            "missing or absent; say it could not be verified from the text.]"
        ))]

    # A single unreadable file must never crash the app — surface a placeholder,
    # keeping the embedded-image note even when the body fails to parse.
    note = _embedded_image_note(lowered, data)
    try:
        if lowered.endswith(".docx"):
            text = note + _read_docx(data)
        elif lowered.endswith(".pdf"):
            text = note + _read_pdf(data)
        elif lowered.endswith((".xlsx", ".xlsm")):
            text = note + _read_xlsx(data)
        elif lowered.endswith(".xlsb"):
            text = note + _read_xlsb(data)
        elif lowered.endswith(".xls"):
            text = note + _read_xls(data)
        elif lowered.endswith(_TEXT_EXTS):
            text = _read_text(data)
            if _looks_binary(text):
                # A renamed binary (a real .xls saved as .csv is a common user
                # mistake) or an undetectable encoding — fail loudly rather
                # than feeding the model garbage that LOOKS like content.
                text = (
                    f"[Could not read {name}: the file decodes to binary "
                    "garbage, not text. It may be a renamed binary file "
                    "(e.g. a real Excel workbook saved with a .csv name) or "
                    "use an unsupported encoding - re-export it as CSV/text "
                    "or upload the original format.]"
                )
        else:
            # Unknown extension: best-effort decode, but never feed the model
            # silent mojibake from an unsupported binary format.
            text = _read_text(data)
            if _looks_binary(text):
                text = (
                    f"[Unsupported binary file format ({name}): text could not "
                    "be extracted. Convert it to docx/pdf/xlsx/csv if its "
                    "content should be graded; do not treat its content as "
                    "missing.]"
                )
    except Exception as exc:
        text = note + f"[Could not read {name}: {exc}]"
    return [ParsedDocument(name, text)]


def _read_member(
    zf: zipfile.ZipFile, info: zipfile.ZipInfo, budget: List[int]
) -> bytes:
    """Decompress one member within the per-file and per-upload byte limits."""
    limit = min(_MAX_MEMBER_BYTES, budget[0])
    too_large = ValueError(
        "too large to open safely inside a zip (limits: "
        f"{_MAX_MEMBER_BYTES >> 20} MB per file, {_MAX_ZIP_TOTAL_BYTES >> 20} "
        "MB per upload) - upload it separately or reduce its size"
    )
    if info.file_size > limit:
        raise too_large
    # The header's file_size is attacker-controlled: the bounded read enforces
    # the limit even if it lies.
    with zf.open(info) as member:
        inner = member.read(limit + 1)
    if len(inner) > limit:
        raise too_large
    budget[0] -= len(inner)
    return inner


def _parse_zip(
    data: bytes, prefix: str = "", depth: int = 1, budget: List[int] | None = None
) -> List[ParsedDocument]:
    # One byte budget per top-level upload, shared by every nested zip.
    budget = [_MAX_ZIP_TOTAL_BYTES] if budget is None else budget
    out: List[ParsedDocument] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            base = info.filename.rsplit("/", 1)[-1]
            if info.filename.startswith(_SKIP_PREFIXES) or base.startswith("."):
                continue
            # Normalise "./" and "../" prefixes instead of dropping the member
            # — nothing is extracted to disk, so the name is only a label.
            # A nested zip's members keep its path as their folder: in LMS
            # exports the student lives in "Name_ID_assignsubmission_file_/",
            # and without it two students' "sub.zip/Report.docx" would merge.
            name = prefix + (re.sub(r"^(\.\.?/)+", "", info.filename) or base)
            try:
                # Reading is inside the try: an encrypted member raises
                # RuntimeError, and one bad file must not sink the batch.
                if not name.lower().endswith(".zip"):
                    out.extend(parse_bytes(name, _read_member(zf, info, budget)))
                elif depth >= _MAX_ZIP_DEPTH:
                    raise ValueError(
                        f"zip nested more than {_MAX_ZIP_DEPTH} levels deep was "
                        "not opened - extract it and upload its files directly"
                    )
                else:
                    inner = _read_member(zf, info, budget)
                    out.extend(_parse_zip(inner, name[:-4] + "/", depth + 1, budget))
            except Exception as exc:
                out.append(ParsedDocument(name, f"[Could not read file: {exc}]"))
    return out


def parse_upload(uploaded_file) -> List[ParsedDocument]:
    """Parse a Streamlit UploadedFile (or anything with .name and .getvalue())."""
    return parse_bytes(uploaded_file.name, uploaded_file.getvalue())


# --- Readability health check ----------------------------------------------
# Scanned PDFs (no text layer), read failures and image-only submissions carry
# little or no gradable text. Grading them would silently produce a low mark
# based on nothing, so the UI surfaces these BEFORE the batch runs.

_ANNOTATION_RES = (
    re.compile(r"\[NOTE: this document contains[^\]]*\]"),
    re.compile(r"\[Image file[^\]]*\]"),
    re.compile(r"\[Could not read[^\]]*\]"),
    re.compile(r"\[Unsupported binary file format[^\]]*\]"),
    re.compile(r"=== File: [^=]*==="),
)


def text_health(text: str) -> str | None:
    """A short problem label when extracted text is unusable for grading.

    Returns None when the text looks gradable. Labels:
      'read-failed'        - a file could not be parsed at all
      'unsupported-format' - binary format we cannot extract
      'no-text'            - nothing readable (typically a scanned PDF or an
                             image-only submission)
      'little-text'        - under ~300 characters of real content
    """
    raw = text or ""
    body = raw
    for pattern in _ANNOTATION_RES:
        body = pattern.sub("", body)
    body = body.strip()
    # Volume alone is not health: mojibake (wrong encoding, renamed binary)
    # easily clears 300 chars but is ungradable — flag it, never bless it.
    if _looks_binary(body):
        return "read-failed"
    if len(body) >= 300:
        return None
    if "[Could not read" in raw:
        return "read-failed"
    if "[Unsupported binary file format" in raw:
        return "unsupported-format"
    if len(body) < 20:
        return "no-text"
    return "little-text"


_FAILURE_REASON_RE = re.compile(r"\[Could not read [^:\]]*:\s*([^\]]+)\]")


def read_failure_reason(text: str) -> str | None:
    """The exception message inside the first '[Could not read ...]' note.

    Surfacing the actual reason in the UI turns "file could not be parsed"
    into something actionable (e.g. "The 'pyxlsb' package is required...").
    """
    match = _FAILURE_REASON_RE.search(text or "")
    if not match:
        return None
    reason = " ".join(match.group(1).split())
    return reason[:300] if reason else None


HEALTH_HINTS = {
    "read-failed": "file could not be parsed - re-export it (e.g. save as a "
                   "new .docx/.pdf) or convert the format",
    "unsupported-format": "unsupported format - convert to docx/pdf/xlsx/csv",
    "no-text": "no readable text - likely a scanned PDF or image-only "
               "submission; run OCR or request a text version",
    "little-text": "very little readable text - check the file before grading",
}


# --- Grouping files into per-student submissions ---------------------------
# A submission often spans several files (a Word report + an Excel workbook).
# Group by student ID in the filename (e.g. s4180007); files without an ID
# group by their normalised filename stem, so "X.docx" pairs with "X.xlsx".

# The lookbehind stops the trailing "s" of ordinary words matching ("Results"
# + a yyyymmdd date would otherwise read as a student ID and merge two
# different students whose filenames share the date).
_STUDENT_ID_RE = re.compile(r"(?<![A-Za-z0-9])[sS](\d{6,8})")


def _group_key(name: str) -> tuple:
    path = name.replace("\\", "/")
    # Search the FULL path: in LMS exports (Moodle/Canvas) the student ID often
    # lives in the per-student folder name, not in the file name itself.
    match = _STUDENT_ID_RE.search(path)
    if match:
        return ("id", match.group(1))
    folder, _, base = path.rpartition("/")
    stem = base.rsplit(".", 1)[0].lower()
    stem = re.sub(r"\s*\(\d+\)\s*$", "", stem)      # drop copy suffixes " (1)"
    stem = re.sub(r"[^a-z0-9]+", "", stem)           # normalise punctuation
    # The parent folder is part of the identity: LMS zips give every student a
    # folder containing generically named files ("report.docx"), and without
    # the folder the whole class would collapse into one submission.
    folder_norm = re.sub(r"[^a-z0-9]+", "", folder.lower())
    key = f"{folder_norm}/{stem}" if folder_norm else (stem or base.lower())
    return ("stem", key)


def _primary_member(members: List[ParsedDocument]) -> ParsedDocument:
    return next(
        (m for m in members if m.name.lower().endswith((".docx", ".pdf"))),
        members[0],
    )


def _merge_members(members: List[ParsedDocument]) -> ParsedDocument:
    """One submission from a group's files: the primary (Word/PDF) file's name
    labels the group and its text comes first, so if a length cap ever trims
    the combined text it is the auxiliary spreadsheets that get shortened, not
    the write-up."""
    if len(members) == 1:
        return members[0]
    primary = _primary_member(members)
    ordered = [primary] + [m for m in members if m is not primary]
    text = "\n\n".join(
        f"=== File: {m.name} ===\n{m.text.strip()}" for m in ordered
    )
    return ParsedDocument(f"{primary.name} [+{len(members) - 1} file(s)]", text)


def _auto_keys(docs: List[ParsedDocument]) -> List[tuple]:
    """The automatic group key for each doc, aligned with the input order."""
    # LMS (Moodle/Canvas) zips give each student their own folder holding a
    # couple of generically named files. For ID-less files, a small folder
    # (<= 4 files) is treated as one student's submission; big folders (a
    # whole-class dump, or "Word/" vs "Excel/" type folders) are not, so a
    # class can never collapse into one group.
    def _folder(name: str) -> str:
        return name.replace("\\", "/").rpartition("/")[0]

    folder_counts: dict = {}
    for doc in docs:
        folder_counts[_folder(doc.name)] = folder_counts.get(_folder(doc.name), 0) + 1

    keys: List[tuple] = []
    for doc in docs:
        key = _group_key(doc.name)
        folder = _folder(doc.name)
        if key[0] == "stem" and folder and folder_counts.get(folder, 0) <= 4:
            key = ("folder", folder)
        keys.append(key)
    return keys


def _partition(docs: List[ParsedDocument], keys: List[tuple]) -> List[List[ParsedDocument]]:
    groups: dict = {}
    order: List[tuple] = []
    for doc, key in zip(docs, keys):
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(doc)
    return [groups[key] for key in order]


def group_by_student(
    docs: List[ParsedDocument], enabled: bool = True
) -> List[ParsedDocument]:
    """Merge files that belong to the same student into one submission."""
    if not enabled:
        return list(docs)
    return [_merge_members(m) for m in _partition(docs, _auto_keys(docs))]


def auto_group_labels(
    docs: List[ParsedDocument], enabled: bool = True
) -> dict[str, str]:
    """A human-editable group label per file name.

    The labels induce the SAME partition as ``group_by_student`` (each auto
    group gets one label, unique across groups), so they can prefill a manual
    grouping editor: the professor merges two groups by simply giving their
    rows the same label. Labels are the primary member's filename stem —
    short, recognisable, and natural to type (e.g. "Team 7").
    """
    keys = (
        _auto_keys(docs) if enabled else [("file", i) for i in range(len(docs))]
    )
    labels: dict[str, str] = {}
    used: dict[str, int] = {}
    for members in _partition(docs, keys):
        primary = _primary_member(members)
        base = primary.name.replace("\\", "/").rpartition("/")[2]
        base = base.rsplit(".", 1)[0].strip() or primary.name
        n = used.get(base.lower(), 0) + 1
        used[base.lower()] = n
        label = base if n == 1 else f"{base} ({n})"
        for m in members:
            labels[m.name] = label
    return labels


def group_by_labels(
    docs: List[ParsedDocument], labels: dict[str, str]
) -> List[ParsedDocument]:
    """Merge files that share the same (professor-assigned) group label.

    Label matching is case-insensitive, so hand-typed "team 7" joins
    "Team 7". A file with an empty/missing label stays its own submission.
    """
    keys: List[tuple] = []
    for i, doc in enumerate(docs):
        label = (labels.get(doc.name) or "").strip()
        keys.append(("label", label.lower()) if label else ("single", i))
    return [_merge_members(m) for m in _partition(docs, keys)]
