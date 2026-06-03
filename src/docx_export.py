"""Generate the downloadable Word (.docx) grade report.

The output replicates the rubric structure for every student: identifier, a
component-by-component table (score / max / professor's notes), the total out
of 100, and the band. A cohort summary page leads the document so the professor
can see the calibrated distribution at a glance.
"""
from __future__ import annotations

import io
from typing import List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .calibration import CalibrationReport
from .models import CalibratedReport


_HEADER_FILL = "2E4057"


def build_grade_document(
    reports: List[CalibratedReport],
    session_name: str,
    calibration: Optional[CalibrationReport] = None,
) -> bytes:
    """Build the .docx and return its bytes (ready for download)."""
    doc = Document()
    _build_styles(doc)

    title = doc.add_heading(f"Grade Report — {session_name}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if calibration is not None:
        _add_cohort_summary(doc, calibration)

    for i, report in enumerate(reports):
        if i > 0 or calibration is not None:
            doc.add_page_break()
        _add_student_section(doc, report)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)


def _add_cohort_summary(doc: Document, calib: CalibrationReport) -> None:
    doc.add_heading("Cohort Summary", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"

    def row(label: str, value: str) -> None:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    row("Reports graded", str(calib.n))
    row("Calibrated mean", f"{calib.final_mean:.2f}%  (raw {calib.raw_mean:.2f}%)")
    row("High Distinction (>=80%)", f"{calib.hd_count}  (cap {calib.hd_cap})")
    row("Distinction (70-79%)", f"{calib.di_count}  (cap {calib.di_cap})")
    row("Fails (<50%)", str(calib.fail_count))
    doc.add_paragraph()


def _add_student_section(doc: Document, report: CalibratedReport) -> None:
    heading = doc.add_heading(report.student_identifier, level=1)
    heading.runs[0].font.color.rgb = RGBColor.from_string(_HEADER_FILL)

    meta = doc.add_paragraph()
    meta.add_run("Source: ").bold = True
    meta.add_run(report.source_name)

    # Component table: Component | Score | Max | Professor's Notes
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for cell, label in zip(hdr, ("Component", "Score", "Max", "Professor's Notes")):
        run = cell.paragraphs[0].add_run(label)
        run.bold = True

    for component in report.components:
        cells = table.add_row().cells
        cells[0].text = component.name
        cells[1].text = f"{component.score:g}"
        cells[2].text = f"{component.max_score:g}"
        notes_cell = cells[3]
        notes_cell.text = ""
        for note in component.notes:
            p = notes_cell.add_paragraph(note, style="List Bullet")
            p.paragraph_format.space_after = Pt(0)

    # Total line
    total_p = doc.add_paragraph()
    total_run = total_p.add_run(
        f"Total: {report.final_total:g} / 100   ·   {report.band}"
    )
    total_run.bold = True
    total_run.font.size = Pt(12)
    total_run.font.color.rgb = RGBColor.from_string(_HEADER_FILL)
