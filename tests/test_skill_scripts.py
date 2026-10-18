"""The skill's CLI scripts (intake -> calibrate -> export), run the way the
skill runs them: as subprocesses whose output is piped, which on Windows means
a cp1252 console and cp1252 as the default file encoding."""
from __future__ import annotations

import dataclasses
import json
import os
import pathlib
import statistics
import subprocess
import sys

import pytest
from docx import Document

from src.calibration import calibrate_scores
from src.config import DEFAULT_CALIBRATION

SCRIPTS = pathlib.Path(__file__).resolve().parent.parent / "skill" / "rubric-grading" / "scripts"
NAME = "Lê Thị Ngọc"
COHORT = {f"s42000{i:02d}": 40.0 + 3 * i for i in range(12)}


def run(script, *args):
    env = {k: v for k, v in os.environ.items() if k != "PYTHONUTF8"}
    env["PYTHONIOENCODING"] = "cp1252"  # a piped console on a Windows machine
    p = subprocess.run(
        [sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True, env=env
    )
    return p.returncode, p.stdout.decode("utf-8", "replace"), p.stderr.decode("utf-8", "replace")


def write_results(d, totals, maxes=(60, 40)):
    d.mkdir(exist_ok=True)
    for sid, total in totals.items():
        comps = [
            {"name": f"C{i}", "max": m, "score": round(total * m / sum(maxes), 1), "notes": ["Hệ số β"]}
            for i, m in enumerate(maxes)
        ]
        r = {"sid": sid, "student_name": NAME, "components": comps, "total": total,
             "overall_comment": "Nguyễn"}
        (d / f"{sid}.json").write_text(json.dumps(r, ensure_ascii=False), encoding="utf-8")


def calibrate(tmp_path, *extra, totals=COHORT, maxes=(60, 40)):
    write_results(tmp_path / "results", totals, maxes)
    out = tmp_path / "finals.json"
    code, stdout, err = run("calibrate.py", "--results", tmp_path / "results", "--out", out, *extra)
    return code, stdout, err, out


@pytest.mark.parametrize("script", ["intake.py", "calibrate.py", "export_docx.py"])
def test_help(script):
    code, out, _ = run(script, "--help")
    assert code == 0 and "usage:" in out


def test_intake_reads_and_writes_unicode(tmp_path):
    src = tmp_path / "in" / f"s4100001_{NAME}_report.docx"
    src.parent.mkdir()
    doc = Document()
    doc.add_paragraph("Hệ số β = 1.2 - Nguyễn")
    doc.save(src)

    code, out, err = run("intake.py", "--work", tmp_path / "work", src)

    assert code == 0, err
    report = tmp_path / "work" / "packs" / "s4100001" / "report.txt"
    assert "Hệ số β = 1.2 - Nguyễn" in report.read_text(encoding="utf-8")
    assert src.name in out


def test_intake_rebuild_drops_files_moved_to_another_student(tmp_path):
    inp = tmp_path / "in"
    inp.mkdir()
    files = {"s4100001_report.txt": "A report", "s4100001_appendix.txt": "B work, misfiled",
             "s4100002_report.txt": "B report"}
    for name, text in files.items():
        (inp / name).write_text(text, encoding="utf-8")
    work = tmp_path / "work"
    assert run("intake.py", "--work", work, *sorted(inp.iterdir()))[0] == 0
    a, b = work / "packs" / "s4100001", work / "packs" / "s4100002"
    assert (a / "report_2.txt").exists()
    (a / "NOTE.txt").write_text("chart image belongs to this student", encoding="utf-8")

    # The professor moves the misfiled appendix to s4100002 by hand (SKILL.md step 2).
    reg_path = work / "registry.json"
    reg = json.loads(reg_path.read_text(encoding="utf-8"))
    key = next(k for k, i in reg["files"].items() if i["orig"] == "s4100001_appendix.txt")
    reg["files"][key]["sid"] = "s4100002"
    reg["students"]["s4100001"]["files"].remove(key)
    reg["students"]["s4100002"]["files"].append(key)
    reg_path.write_text(json.dumps(reg), encoding="utf-8")
    code, _, err = run("intake.py", "--work", work, inp / "s4100001_report.txt")

    assert code == 0, err
    assert sorted(p.name for p in a.iterdir()) == ["MANIFEST.txt", "NOTE.txt", "report.txt"]
    assert (a / "report.txt").read_text(encoding="utf-8") == "A report"
    assert (b / "report_2.txt").read_text(encoding="utf-8") == "B work, misfiled"


def test_intake_refuses_pack_dir_outside_packs(tmp_path):
    src = tmp_path / "s4100001_report.txt"
    src.write_text("A report", encoding="utf-8")
    work = tmp_path / "work"
    assert run("intake.py", "--work", work, src)[0] == 0
    keep = tmp_path / "report.txt"
    keep.write_text("not intake's file", encoding="utf-8")
    reg_path = work / "registry.json"
    reg = json.loads(reg_path.read_text(encoding="utf-8"))
    reg["students"] = {"../..": reg["students"]["s4100001"]}
    reg_path.write_text(json.dumps(reg), encoding="utf-8")

    code, _, err = run("intake.py", "--work", work, src)

    assert code != 0 and "../.." in err
    assert keep.read_text(encoding="utf-8") == "not intake's file"


def test_calibrate_reports_real_caps_and_spread(tmp_path):
    totals = {f"s42{i:05d}": float(i) for i in range(100)}

    code, stdout, err, out = calibrate(tmp_path, "--hd", "29", totals=totals)

    assert code == 0, err
    summary = json.loads(out.read_text(encoding="utf-8"))["summary"]
    assert (summary["hd_cap"], summary["di_cap"]) == (29, 20)
    assert summary["raw_std"] == round(statistics.pstdev(totals.values()), 2)
    assert summary["low_discrimination"] is False
    assert '"hd_cap": 29' in stdout


@pytest.mark.parametrize("flag,value", [("--fixed", "s4209999=50"), ("--no-floor", "s4209999")])
def test_calibrate_rejects_unknown_student_ids(tmp_path, flag, value):
    code, _, err, out = calibrate(tmp_path, flag, value)

    assert code != 0 and "s4209999" in err
    assert not out.exists()


def test_preserve_bands_rejects_typo_and_missing_previous(tmp_path):
    code, _, err, prev = calibrate(tmp_path)
    assert code == 0, err
    out = tmp_path / "lifted.json"
    base = ["calibrate.py", "--results", tmp_path / "results", "--out", out,
            "--preserve-bands", prev, "--mean", "67"]

    code, _, err = run(*base, "--fixed", "s420001=50")
    assert code != 0 and "s420001" in err and not out.exists()

    data = json.loads(prev.read_text(encoding="utf-8"))
    del data["finals"]["s4200005"]
    prev.write_text(json.dumps(data), encoding="utf-8")
    code, _, err = run(*base)
    assert code != 0 and "s4200005" in err and not out.exists()


@pytest.mark.parametrize("others,raw,expected", [(95.0, 2.0, 0.0), (10.0, 99.0, 100.0)])
def test_no_floor_final_stays_within_0_100(tmp_path, others, raw, expected):
    totals = {f"s42000{i:02d}": others + i * 0.1 for i in range(11)}
    totals["s4209999"] = raw

    code, _, err, out = calibrate(tmp_path, "--no-floor", "s4209999", totals=totals)

    assert code == 0, err
    assert json.loads(out.read_text(encoding="utf-8"))["finals"]["s4209999"] == expected


@pytest.mark.parametrize("n,field,frac,expected", [
    (100, "hd_cap_fraction", 0.29, 29),  # 0.29 * 100 == 28.999999999999996
    (180, "di_cap_fraction", 0.35, 63),  # 0.35 * 180 == 62.99999999999999
])
def test_caps_survive_float_rounding(n, field, frac, expected):
    cfg = dataclasses.replace(DEFAULT_CALIBRATION, **{field: frac})
    _, rep = calibrate_scores([float(i % 100) for i in range(n)], cfg)
    assert (rep.hd_cap if field.startswith("hd") else rep.di_cap) == expected


def test_export_shows_real_caps_and_unicode_names(tmp_path):
    res = tmp_path / "results"
    write_results(res, {"s4200001": 70.0, "s4200002": 60.0, "s4200003": 50.0})
    finals = tmp_path / "finals.json"
    finals.write_text(json.dumps({
        "finals": {"s4200001": 82.0, "s4200002": 72.0, "s4200003": 60.0},
        "summary": {"hd": 1, "di": 1, "hd_cap": 2, "di_cap": 4, "shift": 12.0},
    }), encoding="utf-8")
    out = tmp_path / "Bảng điểm.docx"

    code, stdout, err = run("export_docx.py", "--results", res, "--finals", finals,
                            "--title", "Kỳ 1", "--out", out)

    assert code == 0, err
    assert out.name in stdout
    doc = Document(out)
    cells = {r.cells[0].text: r.cells[1].text for r in doc.tables[0].rows}
    assert cells["High Distinction (>=80%)"] == "1  (cap 2)"
    assert cells["Distinction (70-79%)"] == "1  (cap 4)"
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading 1")]
    assert "Le Thi Ngoc (s4200001)" in headings
