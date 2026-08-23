"""Tests for keyboard-only note sanitisation (no smart quotes/dashes/etc.)."""
from __future__ import annotations

from src.docx_export import sanitize_text


def test_user_example_characters():
    before = (
        "Firm data present (≈$600bn assets, $1.8bn Q1 net income, NII +6%, "
        "NIM +11bp, FirstBank acquisition) but description crammed into a single "
        "paragraph with no distinct industry characterisation"
    )
    after = sanitize_text(before)
    assert "≈" not in after and "~$600bn" in after


def test_en_dash_becomes_hyphen():
    s = sanitize_text("No theoretical framing at this stage – concept surfaces later")
    assert "–" not in s
    assert " - concept" in s


def test_smart_quotes_arrows_ellipsis():
    s = sanitize_text("“clear” ‘risk’ hedge → payoff…")
    assert s == '"clear" \'risk\' hedge -> payoff...'


def test_result_is_pure_ascii():
    messy = "–—‘’“”…≈≠≤≥±×→ €600 °C Phạm"
    out = sanitize_text(messy)
    assert out.isascii(), f"non-ASCII left: {out!r}"


def test_transliterates_accents():
    assert "Pham" in sanitize_text("Phạm's report")  # ạ -> a


def test_empty_and_plain_unchanged():
    assert sanitize_text("") == ""
    assert sanitize_text("Weak analysis of market mechanisms") == (
        "Weak analysis of market mechanisms"
    )


def test_document_is_ascii_everywhere():
    """Headings, source line, title, totals and the warning must all be ASCII."""
    import io
    import docx as docx_lib
    from src.calibration import CalibrationReport
    from src.docx_export import build_grade_document
    from src.models import CalibratedReport, ComponentGrade

    rep = CalibratedReport(
        student_identifier="Nguyễn Văn Ăn — ⚠️ (s4100009)",
        source_name="Phạm_bài–làm.docx [+1 file(s)]",
        components=[ComponentGrade(name="Phân tích – đánh giá", score=10,
                                   max_score=15, notes=["ghi chú – ≈5%"])],
        raw_total=70, final_total=68.5,
    )
    calib = CalibrationReport(n=3, raw_mean=70, final_mean=68.5, hd_count=0,
                              di_count=0, fail_count=0, hd_cap=0, di_cap=0,
                              shift=0, raw_std=0.2, distinct_finals=1)
    data = build_grade_document([rep], "Kỳ thi – 2026", calib, graded_by="Claude")
    doc = docx_lib.Document(io.BytesIO(data))
    blob = " ".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    )
    assert blob.isascii(), sorted({ch for ch in blob if ord(ch) > 127})
    assert "Nguyen Van An" in blob        # name transliterated, not dropped
    assert "Phan tich - danh gia" in blob  # Vietnamese d preserved as d


def test_vietnamese_d_stroke_preserved():
    assert sanitize_text("đánh giá Đại học") == "danh gia Dai hoc"
