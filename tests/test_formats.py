"""Format matrix: every submission format must read cleanly or fail loudly,
and image evidence must never be mistaken for absent."""
from __future__ import annotations

import base64
import io
import zipfile

from src.document_parser import (
    ParsedDocument,
    _looks_binary,
    group_by_student,
    parse_bytes,
)

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


# --- CSV / text encodings ---------------------------------------------------

def test_csv_utf8_bom():
    data = b"\xef\xbb\xbf" + "name,score\nAnalysis,60\n".encode("utf-8")
    text = parse_bytes("grades.csv", data)[0].text
    assert text.startswith("name,score")
    assert "﻿" not in text


def test_csv_utf16_excel_export():
    data = "Criterion\tWeight\nAnalysis\t60\n".encode("utf-16")
    text = parse_bytes("export.csv", data)[0].text
    assert "Analysis" in text and "60" in text
    assert "\x00" not in text


def test_csv_cp1252():
    data = "Student,Café résumé\n".encode("cp1252")
    text = parse_bytes("names.csv", data)[0].text
    assert "Café résumé" in text


# --- .xls (legacy) and .xlsb -------------------------------------------------

def test_xls_legacy_reads_rows():
    import xlwt  # test-only writer for legacy xls

    book = xlwt.Workbook()
    sheet = book.add_sheet("Q1")
    sheet.write(0, 0, "GMVP weight")
    sheet.write(0, 1, 0.191)
    buf = io.BytesIO()
    book.save(buf)
    text = parse_bytes("workbook.xls", buf.getvalue())[0].text
    assert "GMVP weight" in text and "0.191" in text


def test_xlsb_with_images_never_silently_garbles():
    # A fabricated .xlsb (zip with xl/media image but an unreadable body):
    # the embedded-image note must survive even though the body fails.
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/media/image1.png", _PNG)
        zf.writestr("xl/workbook.bin", b"\x00\x01\x02 not a real workbook")
    text = parse_bytes("solver.xlsb", buf.getvalue())[0].text
    assert "contains 1 embedded image(s)" in text
    assert "Could not read" in text  # loud failure, not silent mojibake


def test_xlsx_embedded_solver_screenshot_is_noted():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Solver result", 0.555])
    raw = io.BytesIO()
    wb.save(raw)
    # Re-pack with an image under xl/media (how pasted screenshots are stored).
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/media/image1.png", _PNG)
    text = parse_bytes("calcs.xlsx", out.getvalue())[0].text
    assert "contains 1 embedded image(s)" in text
    assert "Solver result" in text  # body still parsed


# --- Standalone images and unknown binaries ----------------------------------

def test_standalone_screenshot_file_gets_note_not_garbage():
    text = parse_bytes("solver_screenshot.png", _PNG)[0].text
    assert "Image file" in text
    assert "could not be verified from the text" in text
    assert "�" not in text


def test_standalone_image_groups_with_student_submission():
    docs = [
        ParsedDocument("s4106125_report.docx", "essay body"),
        *parse_bytes("s4106125_solver.png", _PNG),
    ]
    merged = group_by_student(docs)
    assert len(merged) == 1
    assert "essay body" in merged[0].text
    assert "Image file" in merged[0].text  # grader sees the image EXISTS


def test_unknown_binary_yields_placeholder_not_mojibake():
    text = parse_bytes("old_report.doc", bytes(range(256)) * 40)[0].text
    assert "Unsupported binary file format" in text
    assert "�" not in text


def test_looks_binary_heuristic():
    assert _looks_binary("\x00\x01�" * 100)
    assert not _looks_binary("A normal paragraph of prose, nothing odd here.")


def test_xlsx_formulas_dates_merged_cells():
    import datetime

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Total"
    ws["B1"] = "=SUM(C1:C3)"                  # formula: data_only -> None, skipped
    ws["A2"] = datetime.date(2026, 8, 1)
    ws["B2"] = 0.191
    ws.merge_cells("A3:C3")
    ws["A3"] = "Merged header"
    buf = io.BytesIO()
    wb.save(buf)
    text = parse_bytes("odd.xlsx", buf.getvalue())[0].text
    assert "Total" in text and "Merged header" in text and "0.191" in text
