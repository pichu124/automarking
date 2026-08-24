"""Regression tests for the adversarial-audit fixes.

Each test pins one defect found by the role-based test sweep: crashes on bad
archives, silent mojibake, invented grades, tables that don't add up, boundary
scores surviving their own guard, stale UI state, and estimator drift.
"""
from __future__ import annotations

import dataclasses
import io
import math
import zipfile

import pytest

from src.calibration import calibrate_scores, rescale_components
from src.config import DEFAULT_CALIBRATION as CFG
from src.document_parser import ParsedDocument, group_by_student, parse_bytes, text_health
from src.grading_engine import (
    GradingSession,
    ProviderSpec,
    _is_provider_unhealthy,
    grade_batch,
    grade_batch_with_fallback,
)
from src.models import ComponentGrade, GradingResult, RubricComponent
from src.providers import ModelOutputError


# --- Parser: archives must degrade loudly, never crash ------------------------

def test_corrupt_and_empty_zip_yield_placeholder_not_crash():
    for data in (b"", b"this is not a zip at all" * 10):
        docs = parse_bytes("class_submissions.zip", data)
        assert len(docs) == 1
        assert "[Could not read" in docs[0].text


def test_encrypted_zip_member_yields_placeholder_not_crash():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("ok.txt", "readable essay text here")
        info = zipfile.ZipInfo("secret.docx")
        zf.writestr(info, b"cipher")
    # Mark the member encrypted (flag bit 0) so zf.read raises RuntimeError.
    raw = bytearray(buf.getvalue())
    idx = raw.find(b"secret.docx")
    raw[idx - 24] |= 0x01  # local header general-purpose flag
    docs = parse_bytes("mixed.zip", bytes(raw))
    names = [d.name for d in docs]
    assert "ok.txt" in names
    assert any("[Could not read" in d.text for d in docs if d.name == "secret.docx")


def test_dot_slash_zip_members_are_parsed_not_dropped():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("./s4180007_report.txt", "the essay body " * 30)
        zf.writestr("../evil.txt", "still a student file " * 30)
        zf.writestr(".DS_Store", b"junk")            # real junk stays skipped
        zf.writestr("__MACOSX/x.txt", b"junk")
    docs = parse_bytes("converted.zip", buf.getvalue())
    names = sorted(d.name for d in docs)
    assert names == ["evil.txt", "s4180007_report.txt"]  # normalised, kept


# --- Parser: encodings must never produce silent mojibake ---------------------

def test_utf16_csv_without_bom_decodes_cleanly():
    data = "Criterion,Weight\nAnalysis,60\n".encode("utf-16-le")  # no BOM
    text = parse_bytes("export.csv", data)[0].text
    assert "Analysis" in text and "60" in text
    assert "\x00" not in text
    assert text_health(text * 20) is None


def test_binary_renamed_to_csv_fails_loudly():
    fake_xls = bytes(range(256)) * 40  # a binary blob with a .csv name
    text = parse_bytes("grades.csv", fake_xls)[0].text
    assert "[Could not read" in text
    assert text_health(text) == "read-failed"


def test_mojibake_body_is_never_labelled_healthy():
    garbage = ("\x00\x01\x02�garbled" * 60)  # >300 chars of control soup
    assert text_health(garbage) == "read-failed"


def test_read_failure_reason_is_surfaced():
    from src.document_parser import read_failure_reason

    docs = parse_bytes("Team02_Excel.xlsb", b"\x00\x01 not a workbook" * 50)
    assert text_health(docs[0].text) == "read-failed"
    assert read_failure_reason(docs[0].text) == "File is not a zip file"
    # A missing-dependency failure names the fix directly.
    fake = "[Could not read a.xlsb: The 'pyxlsb' package is required to read .xlsb files. Install it with: pip install pyxlsb]"
    assert "pip install pyxlsb" in read_failure_reason(fake)
    assert read_failure_reason("healthy essay text") is None


def test_manual_grouping_merges_a_teams_differently_named_files():
    from src.document_parser import auto_group_labels, group_by_labels

    # The real case: Team 7's three files have three different name stems, so
    # auto-grouping keeps them apart (11 groups from 7 teams).
    docs = [
        ParsedDocument("TradeNotesTeam 7.csv", "trade notes"),
        ParsedDocument("ClosePosition_Team 7.csv", "close positions"),
        ParsedDocument("Team 7.docx", "the write-up"),
        ParsedDocument("Team 6.docx", "team 6 report"),
        ParsedDocument("Team 6.xlsx", "team 6 workbook"),
    ]
    labels = auto_group_labels(docs)
    # Auto labels induce the auto partition: Team 6's pair shares one label...
    assert labels["Team 6.docx"] == labels["Team 6.xlsx"]
    # ...while Team 7's three files start out split.
    assert len({labels[d.name] for d in docs[:3]}) == 3

    # The professor types the same label on Team 7's rows (case-insensitive).
    labels["TradeNotesTeam 7.csv"] = "team 7"
    labels["ClosePosition_Team 7.csv"] = "Team 7"
    labels["Team 7.docx"] = "Team 7"
    merged = group_by_labels(docs, labels)
    assert len(merged) == 2
    team7 = next(m for m in merged if "Team 7.docx" in m.name)
    assert "[+2 file(s)]" in team7.name          # docx is the primary label
    assert team7.text.index("the write-up") < team7.text.index("trade notes")


def test_group_labels_empty_label_keeps_file_separate():
    from src.document_parser import group_by_labels

    docs = [ParsedDocument("a.docx", "a"), ParsedDocument("b.docx", "b")]
    merged = group_by_labels(docs, {"a.docx": "", "b.docx": ""})
    assert len(merged) == 2


def test_auto_labels_stay_unique_for_identically_named_groups():
    from src.document_parser import auto_group_labels, group_by_labels

    # Two students' folders both hold "report.docx" — their labels must NOT
    # collide, or the manual editor would merge two strangers by default.
    docs = [
        ParsedDocument("Alice Smith_1234/report.docx", "alice"),
        ParsedDocument("Bob Jones_5678/report.docx", "bob"),
    ]
    labels = auto_group_labels(docs)
    assert labels["Alice Smith_1234/report.docx"] != labels["Bob Jones_5678/report.docx"]
    assert len(group_by_labels(docs, labels)) == 2


def test_student_id_never_matches_inside_a_word():
    docs = [
        ParsedDocument("Alice Nguyen - Results20230815.docx", "alice's essay"),
        ParsedDocument("Bob Tran - Results20230815.docx", "bob's essay"),
        ParsedDocument("s4180007_report.docx", "carol's essay"),
        ParsedDocument("folder_s4180007/workbook.xlsx", "carol's workbook"),
    ]
    merged = group_by_student(docs)
    # Alice and Bob stay separate; Carol's two files (real IDs) still merge.
    assert len(merged) == 3


# --- Engine: no invented grades, tables must add up ---------------------------

class _Grader:
    label = "fake"

    def __init__(self, results):
        self.results = results

    def grade_report(self, *, source_name, **_):
        r = self.results[source_name]
        if isinstance(r, Exception):
            raise r
        return r


def _session():
    return GradingSession(
        task_text="t", rubric_text="r",
        components=[RubricComponent(name="C", weight=100.0, description="")],
        reference_samples={},
    )


def _ok(name, score=70.0):
    return GradingResult(
        student_identifier=name,
        components=[ComponentGrade(name="C", score=score, max_score=100.0, notes=["n"])],
        total=score,
    )


def test_empty_components_is_an_error_not_a_credit():
    results = {
        "empty.docx": GradingResult(student_identifier="E", components=[], total=0.0),
        "d0.docx": _ok("d0"), "d1.docx": _ok("d1"), "d2.docx": _ok("d2"),
    }
    docs = [ParsedDocument(n, "x") for n in results]
    res = grade_batch(_Grader(results), _session(), docs)
    assert len(res.reports) == 3                       # empty one excluded...
    assert any("empty.docx" in e for e in res.errors)  # ...and reported loudly
    assert all(r.source_name != "empty.docx" for r in res.reports)


def test_nan_score_becomes_zero_not_full_marks():
    nan_result = GradingResult(
        student_identifier="N",
        components=[
            ComponentGrade(name="A", score=math.nan, max_score=40.0, notes=["n"]),
            ComponentGrade(name="B", score=5.0, max_score=60.0, notes=["n"]),
        ],
        total=math.nan,
    )
    docs = [ParsedDocument("nan.docx", "x")] + [ParsedDocument(f"d{i}.docx", "y") for i in range(3)]
    results = {"nan.docx": nan_result, **{f"d{i}.docx": _ok(f"d{i}") for i in range(3)}}
    res = grade_batch(_Grader(results), _session(), docs)
    nan_report = next(r for r in res.reports if r.source_name == "nan.docx")
    comp_a = next(c for c in nan_report.components if c.name == "A")
    # NaN must be treated as 0 evidence — never "clamped" up to full marks.
    # After rescaling to the calibrated total, A holds at most its share.
    assert comp_a.score <= comp_a.max_score
    assert abs(sum(c.score for c in nan_report.components) - nan_report.final_total) < 0.15


def test_component_maxima_above_100_still_add_up():
    over = GradingResult(
        student_identifier="O",
        components=[
            ComponentGrade(name=f"C{i}", score=30.0, max_score=30.0, notes=["n"])
            for i in range(5)  # maxima sum to 150
        ],
        total=100.0,
    )
    docs = [ParsedDocument("over.docx", "x")] + [ParsedDocument(f"d{i}.docx", "y") for i in range(3)]
    results = {"over.docx": over, **{f"d{i}.docx": _ok(f"d{i}") for i in range(3)}}
    res = grade_batch(_Grader(results), _session(), docs)
    rp = next(r for r in res.reports if r.source_name == "over.docx")
    assert abs(sum(c.score for c in rp.components) - rp.final_total) < 0.15
    assert all(0 <= c.score <= c.max_score for c in rp.components)


def test_context_overflow_keep_chars_parses_api_messages():
    from src.grading_engine import _context_overflow_keep_chars

    anthropic = Exception(
        "Error code: 400 - {'type': 'invalid_request_error', 'message': "
        "'prompt is too long: 1600280 tokens > 1000000 maximum'}"
    )
    keep = _context_overflow_keep_chars(anthropic, 6_400_000)
    assert keep is not None
    # limit/actual * 0.85 => ~53% of the text kept.
    assert 3_000_000 < keep < 3_500_000

    openai_style = Exception(
        "This model's maximum context length is 128000 tokens. However, your "
        "messages resulted in 200000 tokens."
    )
    assert _context_overflow_keep_chars(openai_style, 800_000) is not None

    assert _context_overflow_keep_chars(Exception("invalid json"), 10_000) is None
    # Groq's TPM-style 413 must NOT be treated as an overflow (it falls back).
    assert _context_overflow_keep_chars(
        Exception("Request too large: tokens per minute (TPM) limit 12000"), 10_000
    ) is None


def test_oversized_submission_is_auto_trimmed_and_graded():
    calls = []

    class OverflowGrader:
        label = "primary"

        def grade_report(self, *, source_name, report_text, **_):
            calls.append(len(report_text))
            if len(report_text) > 100_000:
                raise Exception(
                    "Error code: 400 - prompt is too long: "
                    f"{len(report_text) // 4} tokens > 25000 maximum"
                )
            return _ok(source_name)

    huge = ParsedDocument("Team 1.docx [+1 file(s)]", "row,data\n" * 300_000)
    normal = ParsedDocument("Team 2.docx", "a normal report")
    res = grade_batch_with_fallback(
        [ProviderSpec("p", label="primary"), ProviderSpec("q", label="secondary")],
        _session(), [huge, normal],
        builder=lambda spec: OverflowGrader(),
    )
    assert len(res.reports) == 2                 # the huge one still graded
    assert res.errors == []
    assert res.provider_label == "primary"       # no pointless fallback
    assert any("context window" in n for n in res.notes)  # audit trail entry
    assert len(calls) >= 3                       # original + trimmed retries
    assert calls[-1] <= 100_000                  # final attempt fits


def test_context_overflow_that_never_fits_stays_a_per_report_error():
    class AlwaysOverflow:
        label = "primary"

        def grade_report(self, *, source_name, report_text, **_):
            if source_name == "huge.docx":
                raise Exception("prompt is too long: 999999 tokens > 1000 maximum")
            return _ok(source_name)

    docs = [ParsedDocument("huge.docx", "x" * 50_000), ParsedDocument("ok.docx", "fine")]
    res = grade_batch(AlwaysOverflow(), _session(), docs)
    assert len(res.reports) == 1
    assert any("huge.docx" in e for e in res.errors)


def test_model_prose_with_health_words_stays_a_content_error():
    # The model answered ONE report with prose (about a billing system!) instead
    # of JSON — a per-report content error, never a provider-health failure.
    prose_error = ModelOutputError(
        "Model did not return valid JSON. First 200 chars: The submission "
        "provides a strong analysis of the hospital billing workflow and its "
        "rate limit on insurance quota claims"
    )
    assert not _is_provider_unhealthy(prose_error)
    assert not _is_provider_unhealthy(ValueError("something about an api key topic"))

    results = {"d0.txt": _ok("d0"), "d1.txt": prose_error, "d2.txt": _ok("d2")}
    docs = [ParsedDocument(n, "x") for n in results]
    fallback_used = _Grader({n: _ok(n) for n in results})
    res = grade_batch_with_fallback(
        [ProviderSpec("p", label="primary"), ProviderSpec("q", label="secondary")],
        _session(), docs,
        builder=lambda spec: _Grader(results) if spec.provider_key == "p" else fallback_used,
    )
    assert res.provider_label == "primary"      # no cohort-wide switch
    assert res.fell_back_from == []
    assert len(res.reports) == 2
    assert any("d1.txt" in e for e in res.errors)


# --- Calibration: rescaling and boundary guard --------------------------------

def test_rescale_spreads_drift_across_components():
    # A dominant component at its max can't absorb the whole lift alone.
    out = rescale_components([60.0, 0.0, 0.0], [60.0, 20.0, 20.0], 60.0, 88.0)
    assert abs(sum(out) - 88.0) < 0.11
    assert all(0.0 <= s <= m for s, m in zip(out, [60.0, 20.0, 20.0]))

    out = rescale_components([40.0, 5.0, 5.0], [40.0, 30.0, 30.0], 50.0, 85.0)
    assert abs(sum(out) - 85.0) < 0.11


def test_rescale_zero_old_total_is_clamped():
    assert rescale_components([0.0, 0.0], [50.0, 50.0], 0.0, 150.0) == [50.0, 50.0]
    assert rescale_components([0.0, 0.0], [50.0, 50.0], 0.0, -20.0) == [0.0, 0.0]


def test_boundary_guard_never_lands_on_the_boundary_itself():
    # A guard of 1.0 would snap scores onto exactly 69/79 — the scores the
    # feature exists to avoid. The effective guard is floored at 1.5.
    cfg = dataclasses.replace(
        CFG, boundary_guard=1.0,
        target_mean_low=74.5, target_mean_high=75.5,
    )
    raw = [88.0, 78.5, 78.6, 78.4, 78.55]
    final, _ = calibrate_scores(raw, cfg)
    for s in final:
        assert s not in (59.0, 69.0, 79.0), final


def test_mean_lands_in_band_despite_boundary_snapping():
    # Case A from the audit: a cluster snapped at (threshold - guard) makes the
    # mean a step function; the solver must still land inside the band.
    cfg = dataclasses.replace(CFG, target_mean_low=55.0, target_mean_high=56.0)
    raw = [40.0] * 8 + [32.0] * 10
    _, report = calibrate_scores(raw, cfg)
    assert 54.9 <= report.final_mean <= 56.1, report.final_mean

    # Case B: 25 identical raws, band [59, 60], big guard.
    cfg = dataclasses.replace(
        CFG, target_mean_low=59.0, target_mean_high=60.0,
        boundary_guard=2.5, hd_cap_fraction=0.09, di_cap_fraction=0.40,
    )
    raw = [85.6] * 25
    _, report = calibrate_scores(raw, cfg)
    assert 58.9 <= report.final_mean <= 60.1, report.final_mean


# --- Export: non-Latin names must not crash the docx --------------------------

def test_docx_export_survives_non_latin_names():
    from docx import Document as OpenDocument

    from src.docx_export import build_grade_document
    from src.models import CalibratedReport

    reports = [
        CalibratedReport(
            student_identifier=ident, source_name=f"s410000{i}_report.docx",
            components=[ComponentGrade(name="C", score=60.0, max_score=100.0, notes=["ok"])],
            raw_total=60.0, final_total=60.0,
        )
        for i, ident in enumerate(["王小明", "محمد أحمد", "😀🎓", "Nguyễn Văn A"])
    ]
    _, calib = calibrate_scores([60.0] * 4)
    blob = build_grade_document(reports, "Test — 测试", calib, graded_by="Claude")
    doc = OpenDocument(io.BytesIO(blob))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert all(ord(ch) < 128 for ch in text)
    # Identity survives via the source filename fallback.
    assert "s4100000_report.docx" in text


# --- Estimator: count what is actually sent -----------------------------------

def test_estimate_reflects_report_cap_not_raw_size():
    from src import config
    from src.estimator import estimate_batch

    huge = "x" * 5_000_000  # a 5MB document
    est = estimate_batch(
        provider_key="anthropic", provider_label="Claude", model="m",
        static_context_text="ctx " * 100,
        report_texts=[huge], n_components=4,
    )
    capped_tokens = math.ceil(config.MAX_REPORT_CHARS / 4)
    assert est.tokens_per_request_max < capped_tokens + 2000
    assert est.tokens_per_request_max < 100_000  # nowhere near the raw 1.25M


# --- App helpers: upload state must clear when files are removed --------------

class _FakeUpload:
    def __init__(self, name, file_id=None):
        self.name = name
        self.file_id = file_id or name


def test_read_uploads_once_clears_on_remove_but_keeps_manual_edits(monkeypatch):
    import app
    import streamlit as st

    monkeypatch.setattr(
        app, "parse_upload", lambda up: [ParsedDocument(up.name, f"text of {up.name}")]
    )
    for k in ("m_ids", "m_ids_shadow", "m_text"):
        st.session_state.pop(k, None)

    # Upload -> text appears; remove-all -> text clears (the marker-materials case).
    app._read_uploads_once([_FakeUpload("key.docx")], "m_text", "m_ids")
    assert st.session_state["m_text"] == "text of key.docx"
    app._read_uploads_once([], "m_text", "m_ids")
    assert st.session_state["m_text"] == ""

    # Upload -> manual edit -> remove-all: the edit survives (the task case).
    app._read_uploads_once([_FakeUpload("task.docx")], "m_text", "m_ids")
    st.session_state["m_text"] = "my hand-written version"
    app._read_uploads_once([], "m_text", "m_ids")
    assert st.session_state["m_text"] == "my hand-written version"


def test_sample_slot_clears_when_all_files_removed(monkeypatch):
    import app
    import streamlit as st

    monkeypatch.setattr(
        app, "parse_upload", lambda up: [ParsedDocument(up.name, f"text of {up.name}")]
    )
    st.session_state["reference_samples"] = {}
    st.session_state.pop("_ref_HD_ids", None)

    app._read_sample_uploads_once("HD", [_FakeUpload("hd1.docx"), _FakeUpload("hd2.xlsx")])
    assert "=== File: hd1.docx ===" in st.session_state["reference_samples"]["HD"]
    app._read_sample_uploads_once("HD", [])
    assert "HD" not in st.session_state["reference_samples"]


def test_effective_samples_never_serves_stale_profiles(monkeypatch):
    import app
    import streamlit as st

    st.session_state["reference_samples"] = {"HD": "full HD text", "PASS": "full PASS text"}
    st.session_state["use_profiles"] = True
    st.session_state["sample_profiles"] = {"HD": "HD profile", "AVG": "profile of deleted AVG"}
    st.session_state["profile_hashes"] = {
        "HD": app._sample_fingerprint("full HD text"),
        "AVG": app._sample_fingerprint("old avg text"),
    }
    eff = app._effective_samples()
    assert eff == {"HD": "HD profile", "PASS": "full PASS text"}
    # The deleted AVG's profile is gone; the post-compression PASS upload is
    # sent in full; the unchanged HD uses its compact profile.

    # Replacing the HD file invalidates its fingerprint -> full text again.
    st.session_state["reference_samples"]["HD"] = "REPLACED HD text"
    assert app._effective_samples()["HD"] == "REPLACED HD text"
