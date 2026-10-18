"""Policy rules: ties at a cap boundary, pinned marks outside calibration,
rubrics not out of 100, and bare student IDs only after a label."""
from __future__ import annotations

import dataclasses
import io
import itertools
import json
import math
import random
import unicodedata

import pytest
from docx import Document

from src.calibration import CalibrationReport, calibrate_scores
from src.config import DEFAULT_CALIBRATION as CFG
from src.docx_export import build_grade_document
from tests.test_skill_scripts import COHORT, calibrate, run, write_results

BANDS = ((80, "High Distinction"), (70, "Distinction"), (60, "Credit"), (50, "Pass"), (-1, "Fail"))


def band(score):
    return next(name for cut, name in BANDS if score >= cut)


def subdir(tmp_path, name):
    d = tmp_path / name
    d.mkdir()
    return d


# --- A: equal raw totals never split across a cap ----------------------------

@pytest.mark.parametrize("raw,tied", [
    ([85.0, 85.0] + [60.0] * 8, [0, 1]),
    ([60.0] * 8 + [85.0, 85.0], [8, 9]),
])
def test_tie_straddling_hd_cap_is_held_at_distinction(raw, tied):
    final, rep = calibrate_scores(raw)

    assert final[tied[0]] == final[tied[1]] == 78.5
    assert rep.hd_count == 0
    assert rep.cap_ties == [(i, 85.0, "Distinction") for i in tied]


def test_tie_straddling_di_cap_is_held_at_credit():
    raw = [90.0, 80.0, 80.0, 80.0] + [60.0] * 6  # N=10: 1 HD place, 2 DI places

    final, rep = calibrate_scores(raw)

    assert final[1] == final[2] == final[3] == 68.5
    assert final[0] >= 80 and rep.di_count == 0
    assert rep.cap_ties == [(1, 80.0, "Credit"), (2, 80.0, "Credit"), (3, 80.0, "Credit")]


def test_no_cap_ties_when_the_hold_changes_nobody_band():
    _, rep = calibrate_scores([90.0, 80.0, 75.0] + [60.0] * 7)
    assert rep.cap_ties == []


def test_tied_cohorts_are_order_independent_and_keep_the_invariants():
    random.seed(11)
    cohorts = [[float(round(random.gauss(mu, 10))) for _ in range(n)]
               for n in (7, 10, 13, 20, 31) for mu in (60, 72, 85)]
    configs = itertools.product((55.0, 65.5, 75.0), (0, 10, 30), (0, 20, 40), (True, False))
    for mean, hd, di, avoid in configs:
        cfg = dataclasses.replace(
            CFG, target_mean_low=mean - 0.5, target_mean_high=mean + 0.5,
            hd_cap_fraction=hd / 100, di_cap_fraction=di / 100, avoid_band_boundaries=avoid,
        )
        for raw in cohorts:
            raw = [min(100.0, max(0.0, r)) for r in raw]
            n = len(raw)
            final, rep = calibrate_scores(raw, cfg)
            by_raw = {}
            for r, f in zip(raw, final):
                assert by_raw.setdefault(r, f) == f, (cfg, raw, final)
            shuffled = raw[::-1]
            assert [by_raw[r] for r in shuffled] == calibrate_scores(shuffled, cfg)[0]
            order = sorted(range(n), key=lambda i: raw[i], reverse=True)
            assert all(final[a] >= final[b] for a, b in zip(order, order[1:]))
            assert all(0 <= f <= 100 for f in final)
            assert sum(f >= 80 for f in final) <= math.floor(hd / 100 * n + 1e-9)
            assert sum(70 <= f < 80 for f in final) <= rep.di_cap
            for i, r, held in rep.cap_ties:
                assert raw[i] == r and band(final[i]) == held


def _summary_text(calib):
    blob = build_grade_document([], "Tie check", calib)
    doc = Document(io.BytesIO(blob))
    return "\n".join(p.text for p in doc.paragraphs)


def test_docx_summary_lists_cap_ties():
    calib = CalibrationReport(n=10, raw_mean=65, final_mean=65.5, hd_count=0, di_count=2,
                              fail_count=0, hd_cap=1, di_cap=3, shift=2.2,
                              cap_ties=[("Lê Văn A (s4100001)", 85.0, "Distinction"),
                                        ("s4100002", 85.0, "Distinction")])

    text = _summary_text(calib)

    assert ("Tied at the HD/DI cap boundary - held at the lower band: "
            "Le Van A (s4100001) (raw 85, Distinction), s4100002 (raw 85, Distinction)") in text
    calib.cap_ties = []
    assert "Tied at" not in _summary_text(calib)


TIED = {"s4200000": 85.0, "s4200001": 85.0, **{f"s42001{i:02d}": 60.0 + i for i in range(8)}}


def test_calibrate_script_reports_cap_ties(tmp_path):
    code, stdout, err, out = calibrate(tmp_path, totals=TIED)

    assert code == 0, err
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["finals"]["s4200000"] == data["finals"]["s4200001"] < 80
    assert data["summary"]["cap_ties"] == [["s4200000", 85.0, "Distinction"],
                                           ["s4200001", 85.0, "Distinction"]]
    assert "held at the lower band" in stdout and "s4200001" in stdout


def test_export_lists_cap_ties_by_name(tmp_path):
    assert calibrate(tmp_path, totals=TIED)[0] == 0
    out = tmp_path / "grades.docx"

    code, _, err = run("export_docx.py", "--results", tmp_path / "results",
                       "--finals", tmp_path / "finals.json", "--title", "T", "--out", out)

    assert code == 0, err
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert ("held at the lower band: Le Thi Ngoc (s4200000) (raw 85, Distinction), "
            "Le Thi Ngoc (s4200001) (raw 85, Distinction)") in text


# --- B: --fixed marks sit outside calibration ---------------------------------

PINNED = ["s4200000", "s4200001", "s4200002"]


def test_fixed_marks_are_outside_calibration(tmp_path):
    rest = {s: t for s, t in COHORT.items() if s not in PINNED}
    code, _, err, alone = calibrate(subdir(tmp_path, "alone"), totals=rest)
    assert code == 0, err

    code, _, err, out = calibrate(tmp_path, "--fixed", ",".join(f"{s}=35" for s in PINNED))

    assert code == 0, err
    data = json.loads(out.read_text(encoding="utf-8"))
    expected = json.loads(alone.read_text(encoding="utf-8"))["finals"]
    assert data["finals"] == {**expected, **{s: 35.0 for s in PINNED}}
    assert 65.0 <= data["summary"]["final_mean"] <= 66.0
    assert data["summary"]["fixed"] == PINNED


def test_preserve_bands_targets_the_mean_of_the_non_fixed(tmp_path):
    code, _, err, prev = calibrate(tmp_path)
    assert code == 0, err
    out = tmp_path / "lifted.json"

    code, _, err = run("calibrate.py", "--results", tmp_path / "results", "--out", out,
                       "--preserve-bands", prev, "--mean", "67",
                       "--fixed", ",".join(f"{s}=35" for s in PINNED))

    assert code == 0, err
    finals = json.loads(out.read_text(encoding="utf-8"))["finals"]
    others = [v for s, v in finals.items() if s not in PINNED]
    assert abs(sum(others) / len(others) - 67) <= 0.1
    assert all(finals[s] == 35.0 for s in PINNED)


# --- C: rubrics not out of 100 ------------------------------------------------

def test_calibrate_converts_a_40_point_rubric_to_percent(tmp_path):
    points = {sid: 16.0 + 1.25 * i for i, sid in enumerate(COHORT)}  # /40, exact in binary
    percent = {sid: p * 2.5 for sid, p in points.items()}

    code, stdout, err, out40 = calibrate(subdir(tmp_path, "p40"), totals=points, maxes=(25, 15))
    assert code == 0, err
    code, _, err, out100 = calibrate(subdir(tmp_path, "p100"), totals=percent)
    assert code == 0, err

    a, b = (json.loads(p.read_text(encoding="utf-8")) for p in (out40, out100))
    assert a == b
    assert "/40" in stdout


def test_export_shows_marks_on_the_rubric_scale(tmp_path):
    res = tmp_path / "results"
    write_results(res, {"s4200001": 30.0, "s4200002": 22.0}, maxes=(25, 15))
    write_results(res, {"s4200003": 50.0})
    finals = tmp_path / "finals.json"
    finals.write_text(json.dumps({"finals": {"s4200001": 77.2, "s4200002": 62.0, "s4200003": 82.0},
                                  "summary": {}}), encoding="utf-8")
    out = tmp_path / "grades.docx"

    code, _, err = run("export_docx.py", "--results", res, "--finals", finals,
                       "--title", "T", "--out", out)

    assert code == 0, err
    doc = Document(out)
    totals = [p.text for p in doc.paragraphs if p.text.startswith("Total:")]
    assert totals == ["Total: 82 / 100 - High Distinction",
                      "Total: 30.9 / 40 (77.2%) - Distinction",
                      "Total: 24.8 / 40 (62%) - Credit"]
    for table, shown in zip(doc.tables[2:], (30.9, 24.8)):
        rows = table.rows[1:]
        assert [r.cells[2].text for r in rows] == ["25", "15"]
        assert abs(sum(float(r.cells[1].text) for r in rows) - shown) < 0.05


# --- D: bare student IDs only after a label -----------------------------------

def _intake_ids(tmp_path, text):
    src = tmp_path / "report_alpha.txt"
    src.write_text(text, encoding="utf-8")
    code, stdout, err = run("intake.py", "--work", tmp_path / "work", src)
    assert code == 0, err
    reg = json.loads((tmp_path / "work" / "registry.json").read_text(encoding="utf-8"))
    return sorted(reg["students"]), stdout


def test_bare_number_without_label_is_not_a_student_id(tmp_path):
    ids, stdout = _intake_ids(tmp_path, "Portfolio plan\nInitial capital AUD 3000000 in equities.")
    assert ids == [] and "UNMATCHED: report_alpha.txt" in stdout


@pytest.mark.parametrize("label", [
    "Student ID:", "student id -", "Student No.", "Student number", "ID:", "ID No:",
    "Mã SV:", "MSSV", "Mã số sinh viên:", unicodedata.normalize("NFD", "Mã số sinh viên:"),
])
def test_bare_number_after_a_label_is_the_student_id(tmp_path, label):
    ids, _ = _intake_ids(tmp_path, f"Initial capital AUD 3000000\n{label} 4123456\n")
    assert ids == ["s4123456"]


def test_s_prefixed_id_still_found_anywhere(tmp_path):
    ids, _ = _intake_ids(tmp_path, "capital 3000000 " + "x " * 3000 + "by s4123456")
    assert ids == ["s4123456"]


# --- A (app): tie cases are listed for the professor ----------------------------

def test_app_lists_students_held_at_the_lower_band(monkeypatch):
    for key in ("ANTHROPIC_API_KEY", "OPENROUTER_API_KEY"):
        monkeypatch.setenv(key, "dummy")
    import app
    from src.calibration import CalibrationConfig, calibrate_scores

    raw = [85, 85] + [60] * 8
    report = calibrate_scores(raw, CalibrationConfig(hd_cap_fraction=0.1))[1]
    assert report.cap_ties, "fixture must straddle the HD cap"
    message = app._cap_ties_message(report)
    assert "lower band" in message and "85" in message
    assert app._cap_ties_message(calibrate_scores([90, 70, 60], CalibrationConfig())[1]) is None
