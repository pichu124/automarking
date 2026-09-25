"""Write the calibrated cohort to a Word grade report (same layout as the app).

Usage:
    python export_docx.py --results RESULTS_DIR --finals finals.json \
        --title "COURSE Assessment N" --out grades.docx [--names s123="Full Name",...]

Per student: heading (name + ID), component table (Component | Score | Max |
Professor's Notes) with component scores rescaled so they add up to the final
total, the total with its band, and an "Assignment Comments" paragraph. A
cohort summary opens the document. All text is sanitised to keyboard ASCII.
"""
from __future__ import annotations

import argparse
import copy
import glob
import io
import json
import os
import statistics
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from grading_lib.calibration import CalibrationReport, rescale_components  # noqa: E402
from grading_lib.docx_export import build_grade_document, sanitize_text  # noqa: E402
from grading_lib.models import CalibratedReport, ComponentGrade  # noqa: E402


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", required=True)
    ap.add_argument("--finals", required=True)
    ap.add_argument("--title", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--names", default="", help='s123="Full Name",s456="Other"')
    ap.add_argument("--graded-by", default="Claude")
    a = ap.parse_args()

    names = {}
    for part in filter(None, a.names.split(",")):
        k, v = part.split("=", 1)
        names[k.strip()] = v.strip().strip('"')
    fj = json.load(open(a.finals))
    finals, summ = fj["finals"], fj.get("summary", {})

    reports, comments = [], {}
    for p in sorted(glob.glob(os.path.join(a.results, "*.json"))):
        r = json.load(open(p))
        sid, f = r["sid"], float(finals[r["sid"]])
        nm = names.get(sid) or r.get("student_name") or ""
        ident = f"{nm} ({sid})" if nm else sid
        comps = r["components"]
        new = rescale_components(
            [c["score"] for c in comps], [c["max"] for c in comps], sum(c["score"] for c in comps), f
        )
        reports.append(CalibratedReport(
            student_identifier=ident, source_name=sid,
            components=[ComponentGrade(name=c["name"], score=s, max_score=c["max"], notes=c.get("notes") or ["-"])
                        for c, s in zip(comps, new)],
            raw_total=float(r["total"]), final_total=f,
        ))
        comments[sanitize_text(ident)] = r.get("overall_comment", "")

    vals = [r.final_total for r in reports]
    raws = [r.raw_total for r in reports]
    calib = CalibrationReport(
        n=len(vals), raw_mean=round(statistics.mean(raws), 2), final_mean=round(statistics.mean(vals), 2),
        hd_count=sum(v >= 80 for v in vals), di_count=sum(70 <= v < 80 for v in vals),
        fail_count=sum(v < 50 for v in vals), hd_cap=summ.get("hd", 0), di_cap=summ.get("di", 0),
        shift=summ.get("shift", 0.0), raw_std=round(statistics.pstdev(raws), 2) if len(raws) > 1 else 0.0,
        distinct_finals=len(set(vals)), di_cap_widened=bool(summ.get("di_cap_widened")),
    )
    reports.sort(key=lambda x: -x.final_total)
    blob = build_grade_document(reports, a.title, calib, graded_by=a.graded_by)

    doc = Document(io.BytesIO(blob))
    cur = None
    for p in list(doc.paragraphs):
        if p.style.name.startswith("Heading 1") and p.text in comments:
            cur = p.text
        elif p.text.startswith("Total:") and cur and comments.get(cur):
            node = copy.deepcopy(p._p)
            p._p.addnext(node)
            q = Paragraph(node, p._parent)
            for run in q.runs[1:]:
                run._r.getparent().remove(run._r)
            q.runs[0].text = "Assignment Comments: " + sanitize_text(comments[cur])
            q.runs[0].bold = False
            q.runs[0].font.size = Pt(10.5)
            q.runs[0].font.color.rgb = RGBColor(0, 0, 0)
            cur = None
    doc.save(a.out)
    bad = sum(1 for p in Document(a.out).paragraphs for ch in p.text if ord(ch) > 127)
    print(f"saved {a.out}: {len(reports)} students, mean {calib.final_mean}, non-ASCII chars {bad}")


if __name__ == "__main__":
    main()
